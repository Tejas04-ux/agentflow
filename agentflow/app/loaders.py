"""Extract text from uploaded files. Returns a list of (page_number | None, text)."""
from __future__ import annotations

import io
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFileType(ValueError):
    pass


def load_document(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED))}"
        )
    if ext == ".pdf":
        return _load_pdf(data)
    if ext == ".docx":
        return _load_docx(data)
    return [(None, data.decode("utf-8", errors="replace"))]


def _load_pdf(data: bytes) -> list[tuple[int | None, str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append((i, text))
    return pages


def _load_docx(data: bytes) -> list[tuple[int | None, str]]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [(None, "\n\n".join(parts))]
